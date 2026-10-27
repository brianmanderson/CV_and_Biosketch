"""Build NSF_Biosketch_SciENcv_Worksheet.docx (python-docx).

Re-run after editing the data below:  python build_worksheet.py
Then re-export the PDF with Word (see README.md).
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "NSF_Biosketch_SciENcv_Worksheet.docx"
PREPARED = "27 October 2026"
FONT = "Times New Roman"
FLAG_RE = re.compile(r"(\[CONFIRM[^\]]*\]|\[DOI NEEDED[^\]]*\])")

# ----------------------------------------------------------------------------- helpers

def set_font(run, bold=None, italic=None, size=11):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def add_text(par, text, bold=False, italic=False, size=11):
    """Append text to a paragraph; [CONFIRM ...] / [DOI NEEDED ...] tokens get a yellow highlight."""
    for piece in FLAG_RE.split(text):
        if not piece:
            continue
        run = par.add_run(piece)
        set_font(run, bold=bold, italic=italic, size=size)
        if FLAG_RE.fullmatch(piece):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.font.bold = True
    return par


def para(doc, text="", bold=False, italic=False, size=11, after=6, before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if align:
        p.alignment = align
    if text:
        add_text(p, text, bold=bold, italic=italic, size=size)
    return p


def rich(doc, parts, after=6):
    """parts = [(text, bold, italic), ...] in one paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    for text, b, i in parts:
        add_text(p, text, bold=b, italic=i)
    return p


def bullet(doc, text, after=3):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(after)
    add_text(p, text)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    add_text(p, text, bold=True, size=14 if level == 1 else 12)
    return p


def shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def table(doc, headers, rows, widths, size=11):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblpr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    # header row: bold, shaded, repeated on each page
    trpr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    trpr.append(hdr)
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        add_text(cell.paragraphs[0], h, bold=True, size=size)
        shade(cell, "D9D9D9")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            lines = val if isinstance(val, list) else [val]
            p = cells[i].paragraphs[0]
            for j, line in enumerate(lines):
                if j:
                    p = cells[i].add_paragraph()
                add_text(p, line, size=size)
    for row in t.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
            for p in row.cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
    for i, w in enumerate(widths):
        t.columns[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def style_fonts(doc):
    for name, size, bold in (("Normal", 11, None), ("Heading 1", 14, True), ("Heading 2", 12, True),
                             ("List Bullet", 11, None), ("Title", 18, True)):
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        if bold is not None:
            st.font.bold = bold
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), FONT)
        # drop theme font references so Times New Roman wins
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]


def footer(doc):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, f"NSF Biographical Sketch worksheet, B. M. Anderson, prepared {PREPARED}. Proposal for review, not an official document. Page ", size=9)
    run = p.add_run()
    set_font(run, size=9)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._element.append(el)


# ----------------------------------------------------------------------------- data

UCSD = "University of California, San Diego"
UCSD_DEPT = "Department of Radiation Medicine and Applied Sciences"
UCSD_LOC = "San Diego, California, United States"
UNC = "University of North Carolina at Chapel Hill"
UNC_LOC = "Chapel Hill, North Carolina, United States"
MDA = "The University of Texas MD Anderson Cancer Center"
UTGSBS = "The University of Texas MD Anderson Cancer Center UTHealth Houston Graduate School of Biomedical Sciences"
HOU = "Houston, Texas, United States"
GT = "Georgia Institute of Technology"
ATL = "Atlanta, Georgia, United States"

IDENTIFYING = [
    ["Name (Last, First Middle)", "Anderson, Brian Mark", "CV header; SciENcv biosketch certified 26 Feb 2026."],
    ["Persistent identifier (ORCID iD)", "https://orcid.org/0000-0002-2748-1444",
     "CLAUDE.md; SciENcv PDF 02/2026. SciENcv fills this from the ORCID account linked to the NCBI login."],
    ["Position title", "Clinical Assistant Professor of Medical Physics [CONFIRM wording]",
     ["CV header and NIH biosketch say “Clinical Assistant Professor”; the SciENcv biosketch certified 02/2026 says “Assistant Professor, Medical Physics”.",
      "Pick one wording and use it in all four NSF documents (biosketch, Current & Pending, COA, Synergistic Activities)."]],
    ["Organization (primary)", UCSD, "CV header."],
    ["Location (City, State/Province, Country)", "San Diego, California, United States",
     "SciENcv uses drop-downs for state and country; type the city."],
]

PREP_HEAD = ["Organization", "Location", "Degree / entry type", "Start (MM/YYYY)", "Completion (MM/YYYY)", "Field of study", "Source / notes"]
PREP = [
    [UCSD, UCSD_LOC, "Residency, Medical Physics (no degree; the 02/2026 SciENcv entry used “Resident”) [CONFIRM entry type offered by the SciENcv drop-down]",
     "07/2021", "06/2023 [CONFIRM: NIH biosketch table says 7/2023; SciENcv 02/2026 says 06/2023]",
     "Therapeutic Medical Physics (SciENcv 02/2026 used “Therapy Physics”)",
     ["CV Education: 2021 – 2023 Residency, UCSD.", "NSF asks for all postdoctoral/fellowship training; the residency is Brian’s training entry."]],
    [UTGSBS, HOU, "Ph.D.", "08/2017", "06/2021 [CONFIRM: NIH biosketch table says 5/2021; SciENcv 02/2026 says 06/2021]",
     "Medical Physics [CONFIRM: NIH biosketch table lists “Diagnostic Imaging Medical Physics” for the Ph.D. row]",
     ["CV Education: 2017 – 2021 Ph.D., Medical Physics, MD Anderson.", "Dissertation: Improving Treatment of Local Liver Ablation Therapy with Deep Learning and Biomechanical Modeling."]],
    [UTGSBS, HOU, "S.M.S. (Master of Science)", "08/2015", "07/2017",
     "Medical Physics [CONFIRM: NIH biosketch table lists “Therapeutic Medical Physics” for the S.M.S. row]",
     "CV Education: 2015 – 2017 S.M.S, Medical Physics, MD Anderson; NIH biosketch table 7/2017."],
    [GT, ATL, "B.S.", "08/2010", "05/2015", "Nuclear and Radiological Engineering",
     ["CV Education: 2010 – 2015 B.S, Nuclear Engineering, Highest Honors.", "NIH biosketch table: BS, 5/2015, Nuclear and Radiological Engineering."]],
]

APPT_HEAD = ["Start", "End", "Title", "Organization", "Department", "Location", "Source / notes"]
APPT = [
    ["2025", "present", "Clinical Assistant Professor (of Medical Physics)", UCSD, UCSD_DEPT, UCSD_LOC,
     ["CV header; NIH biosketch Positions and Employment.", "NIH biosketch names the department “Radiation Oncology”; the CV service section uses “Radiation Medicine & Applied Sciences”, which is the UCSD department name."]],
    ["2025", "present", "Adjunct Assistant Professor", UNC, "Department of Radiation Oncology [CONFIRM department]", UNC_LOC,
     ["NIH biosketch Positions and Employment; SciENcv 02/2026.", "OUTSIDE THE PRIMARY ORGANIZATION: must be listed (see rule above). Do not omit."]],
    ["2025", "present", "Chair of Technical Innovation and Implementation", UCSD, UCSD_DEPT, UCSD_LOC,
     "CV Professional Service (UCSD Departmental Leadership); NIH biosketch Positions and Scientific Appointments. Titled departmental position: counts as an appointment."],
    ["2025", "present", "Technical Service Chief, Central Nervous System", UCSD, UCSD_DEPT, UCSD_LOC,
     "CV Professional Service; NIH biosketch. Titled departmental position: counts as an appointment."],
    ["2023", "present", "Associate Editor", "Medical Physics (journal of the American Association of Physicists in Medicine)", "N/A",
     "Alexandria, Virginia, United States [CONFIRM location: AAPM headquarters]",
     ["CV Professional Service (Medical Physics journal); NIH biosketch “Associate Editor: The International Journal of Medical Physics Research and Practice”.", "Titled editorial appointment, unpaid: NSF counts it."]],
    ["2023", "2025", "Clinical Assistant Professor", UNC, "Department of Radiation Oncology", UNC_LOC,
     ["NIH biosketch Positions and Employment; SciENcv 02/2026 (listed as “Assistant Professor, Medical Physics”).", "Outside the primary organization and within three years of submission: must be listed."]],
    ["2021", "2023", "Clinical Resident, Medical Physics", UCSD, UCSD_DEPT, UCSD_LOC,
     ["NIH biosketch Positions and Employment; CV Education.", "Also appears under Professional Preparation. The 02/2026 SciENcv biosketch listed it only there; listing it in both places is allowed. [CONFIRM keep here]"]],
    ["2015", "2021", "Graduate Student Researcher", MDA, "Department of Imaging Physics [CONFIRM department]", HOU,
     "NIH biosketch Positions and Employment (“Graduate Student Researcher, MDA Cancer Center, Houston, TX”)."],
    ["2014", "2014", "AAPM Summer Undergraduate Fellow", "St. Joseph Hospital", "Radiation Oncology [CONFIRM department]", "Eureka, California, United States",
     "NIH biosketch Positions and Employment (“2014 AAPM Fellowship, St. Joseph’s Hospital, Eureka, CA”); CV Grants and Fellowships 2014."],
    ["2013", "2014", "Research Assistant", GT, "Nuclear and Radiological Engineering [CONFIRM department]", ATL,
     "NIH biosketch Positions and Employment (“2013-2014 Research Assistant, Georgia Institute of Technology”)."],
]

OPT_HEAD = ["Start", "End", "Title / role", "Organization", "Recommendation"]
OPT = [
    ["2026", "2027", "Search Committee Member (medical physicist search)", UCSD, "Omit: internal committee duty, not a titled appointment. Better suited to Synergistic Activities if at all."],
    ["2026", "present", "Member, QUANTEC 2 Working Group (Q2WG)", "American Association of Physicists in Medicine (AAPM)", "Brian’s call. Committee memberships are usually not listed as appointments; they fit Synergistic Activities. [CONFIRM]"],
    ["2024", "present", "Member, Machine Intelligence Subcommittee", "AAPM", "Same as above. [CONFIRM]"],
    ["2023", "present", "Member, Medical Physics Board of Associate Editors", "AAPM", "Redundant with the Associate Editor appointment already listed; omit."],
    ["2021", "present", "Member, Task Group 263 (Standardizing Nomenclatures in Radiation Oncology)", "AAPM",
     "Same as above. Note the CV gives 2021 – Present in Professional Service but 2021–2025 under Major Contributions. [CONFIRM end date]"],
    ["2021", "2023", "Chief Resident", UCSD, "Titled role (CV Clinical Projects). Optional; the 02/2026 SciENcv omitted it. [CONFIRM]"],
    ["2020", "2021", "Ad hoc Associate Editor", "Medical Physics (AAPM)", "Titled editorial role (CV Professional Service). Optional. [CONFIRM]"],
    ["2019", "present", "Manuscript Reviewer", "Medical Physics (AAPM)", "Omit: reviewing is not a titled appointment."],
    ["2015 / 2017", "present", "Member, AAPM; Member, Medical Physicists for World Benefit", "AAPM; MPWB", "Omit: society memberships are not appointments."],
]

PROD_HEAD = ["#", "Authors (plain text)", "Title", "Venue, volume(issue), pages", "Date", "DOI / URL", "Notes"]
RELATED = [
    ["A1", "Anderson BM, Wahid KA, Brock KK",
     "Simple Python Module for Conversions Between DICOM Images and Radiation Therapy Structures, Masks, and Prediction Arrays",
     "Practical Radiation Oncology, 11(3), 226–229", "05/2021",
     ["https://doi.org/10.1016/j.prro.2021.02.003", "PMCID PMC8102371"],
     ["In PubMed: should appear in My Bibliography / ORCID.", "CV lists the submitted title (“Simple Python Module for Dicom and RT …”) and 02/2021; use the published title."]],
    ["A2", "Anderson BM",
     "DicomRTTool: open-source Python library converting DICOM images, RT structure sets and dose grids to NumPy/SimpleITK arrays, and model predictions back to DICOM-RT structures (software)",
     "GitHub repository brianmanderson/Dicom_RT_and_Images_to_Mask; PyPI package DicomRTTool (first release 09/2020; current v6.2.0, 07/2026)", "09/2020 (first release)",
     ["https://github.com/brianmanderson/Dicom_RT_and_Images_to_Mask", "https://pypi.org/project/DicomRTTool/"],
     ["CV Major Contributions. Software is an allowed NSF product (citable, accessible). Not in ORCID: enter manually in SciENcv using the non-journal citation type.", "Checked 27 Oct 2026: 108 GitHub stars, 32 forks. Usage figures are optional context, not part of the citation."]],
    ["A3", "Anderson BM, Padilla L, Ryckman JM, Covington E, Hong DS, Woods K, Katz MS, Zuhour R, Estes C, Moore KL, Bojechko C",
     "Open RT Structures: A Solution for TG-263 Accessibility",
     "International Journal of Radiation Oncology, Biology, Physics, 118(3), 859–863", "03/2024",
     "https://doi.org/10.1016/j.ijrobp.2023.09.041",
     ["Listed in the 02/2026 SciENcv biosketch with this citation.", "CV gives 06/2024 and omits Zuhour R; Crossref and the 02/2026 SciENcv give the March 2024 issue and 11 authors (CV fix is a separate task)."]],
    ["A4", "Anderson BM, Lin EY, Cardenas CE, Gress DA, Erwin WD, Odisio BC, Koay EJ, Brock KK",
     "Automated Contouring of Contrast and Noncontrast Computed Tomography Liver Images With Fully Convolutional Networks",
     "Advances in Radiation Oncology, 6(1), 100464", "01/2021 (online 05/2020)",
     ["https://doi.org/10.1016/j.adro.2020.04.023", "PMCID PMC7807136"],
     "In PubMed. CV lists 05/2020 (online date); SciENcv 02/2026 used Jan–Feb 2021."],
    ["A5", "Anderson B, Repka M, Gao Y, Byrd E, McGurk R, Das S, Dooley J, Arcot R, Bhattacharya S, Adapa K, Wright J, Marks LB, Mazur LM",
     "Using a Deep Learning System to Delineate the Elective Pelvic Nodal Volumes in Patients with Prostate Cancer: A Tool Aimed to Facilitate Peer Review",
     "Advances in Radiation Oncology, article 102219; volume/issue [CONFIRM: not yet assigned in Crossref on 27 Oct 2026]", "09/2026 (Crossref); CV lists 07/2026",
     "https://doi.org/10.1016/j.adro.2026.102219",
     ["Absent from ORCID (2022+ gap): enter manually or push to ORCID first.", "CV author list omits Dooley J (Crossref lists 13 authors)."]],
]
OTHER = [
    ["B1", "Odisio BC, Albuquerque J, Lin YM, Anderson BM, O’Connor CS, Rigaud B, et al. (30 authors)",
     "Software-based versus visual assessment of the minimal ablative margin in patients with liver tumours undergoing percutaneous thermal ablation (COVER-ALL): a randomised phase 2 trial",
     "The Lancet Gastroenterology & Hepatology, 10(5), 442–451", "05/2025",
     "https://doi.org/10.1016/S2468-1253(25)00024-X",
     ["Phase 2 RCT (NCT04083378) built on Brian’s dissertation registration and segmentation tools (CV Major Contributions).", "Published title reads “randomised phase 2 trial”; the CV adds “controlled”. Crossref spells the 16th author Shah KY (CV: Shan)."]],
    ["B2", "Anderson BM, Lin YM, Lin EY, Cazoulat G, Gupta S, Jones AK, Odisio BC, Brock KK",
     "A novel use of biomechanical model-based deformable image registration (DIR) for assessing colorectal liver metastases ablation outcomes",
     "Medical Physics, 48(10), 6226–6236", "10/2021 (online 08/2021)",
     ["https://doi.org/10.1002/mp.15147", "PMCID PMC9380122"],
     "In PubMed; listed in the 02/2026 SciENcv biosketch. CV lists 07/2021."],
    ["B3", "Anderson B, Moore L, Bojechko C",
     "Rapid in vivo EPID image prediction using a combination of analytically calculated attenuation and AI predicted scatter",
     "Medical Physics, 52(2), 1058–1069", "02/2025 (online 11/2024)",
     "https://doi.org/10.1002/mp.17549",
     "CV Papers 02/2025. Likely absent from ORCID: enter manually."],
    ["B4", "Anderson BM, Bojechko C",
     "DICOM attribute manipulation tool: Easily change frame of reference, series instance, and SOP instance UID",
     "Journal of Applied Clinical Medical Physics, 26(6), e70104", "06/2025 (online 04/2025)",
     "https://doi.org/10.1002/acm2.70104",
     ["Listed in the 02/2026 SciENcv biosketch.", "CV title says “Study Instance UID”; the published title says “SOP instance UID”."]],
    ["B5", "Anderson BM",
     "reconstruct_patches operation and ReconstructPatches2D / ReconstructPatches3D layers for Keras 3: the inverse of extract_patches, reassembling patch-wise model predictions into full-resolution images and volumes (software contribution)",
     "keras-team/keras (GitHub), pull request #22984, merged 16 July 2026; follow-ups #23267 (merged 20 July 2026) and #23319 (merged 4 September 2026)", "07/2026",
     "https://github.com/keras-team/keras/pull/22984",
     ["CV Major Contributions. Enter manually (non-journal type).", "First tagged Keras release that ships it: [CONFIRM] (release notes through v3.15.1 of 29 July 2026 do not name it). Add the keras.io API page URL once it is live, if Brian prefers a docs link to the PR."]],
]

BENCH_HEAD = ["#", "Citation (plain text)", "DOI / URL", "Notes"]
BENCH = [
    ["C1", "Anderson BM, Rigaud B, Lin YM, Jones AK, Kang HC, Odisio BC, Brock KK. Automated segmentation of colorectal liver metastasis and liver ablation on contrast-enhanced CT images. Frontiers in Oncology, 12, 886517, 08/2022.",
     ["https://doi.org/10.3389/fonc.2022.886517", "PMCID PMC9403767"],
     "Was related product #2 in the 02/2026 SciENcv biosketch. Journal is Frontiers in Oncology (Radiation Oncology section); the CV says “Frontiers in Radiation Oncology”."],
    ["C2", "Rigaud B, Anderson BM, Yu ZH, Gobeli M, Cazoulat G, Söderberg J, et al. Automatic Segmentation Using Deep Learning to Enable Online Dose Optimization During Adaptive Radiation Therapy of Cervical Cancer. International Journal of Radiation Oncology, Biology, Physics, 109(4), 1096–1110, 03/2021.",
     "https://doi.org/10.1016/j.ijrobp.2020.10.038",
     "CV lists 10/2020 with the abstract’s title. Do not cite 10.1016/j.ijrobp.2020.07.2573: that is the ASTRO 2020 abstract, not the paper."],
    ["C3", "Cardenas CE, Yang J, Anderson BM, Court LE, Brock KK. Advances in Auto-Segmentation. Seminars in Radiation Oncology, 29(3), 185–197, 07/2019.",
     "https://doi.org/10.1016/j.semradonc.2019.02.001", "Review article; good “other significant” candidate if a segmentation review is wanted."],
    ["C4", "Lin YM, Paolucci I, O’Connor CS, Anderson BM, Rigaud B, Fellman BM, Jones KA, Brock KK, Odisio BC. Ablative Margins of Colorectal Liver Metastases Using Deformable CT Image Registration and Autosegmentation. Radiology, 307(2), e221373, 04/2023.",
     "https://doi.org/10.1148/radiol.221373", "Clinical validation of the dissertation tools."],
    ["C5", "Covington EL, Suresh K, Anderson BM, Barker M, Dess K, Price JG, et al. Perceptions on and roadblocks to implementation of standardized nomenclature in radiation oncology: A survey from TG-263U1. Journal of Applied Clinical Medical Physics, 25(6), e14359, 06/2024.",
     "https://doi.org/10.1002/acm2.14359", "Pairs with Open RT Structures for a standardization-themed proposal."],
    ["C6", "Floyd W, Wahid KA, Anderson B, Knäusl B, Moreno AC. Large language models for scalable standardization in radiation oncology. Physics and Imaging in Radiation Oncology, 38, 100960, 03/2026.",
     "https://doi.org/10.1016/j.phro.2026.100960", "Absent from ORCID; enter manually if used."],
    ["C7", "Gao Y, Wells D, Bhattacharya S, Repka M, Lavrova L, Anderson B, et al. Deep Learning-Based Pelvic Vessel Auto-Segmentation for Standardized Lymph Node Delineation in Prostate Cancer Radiotherapy. Practical Radiation Oncology, 09/2026; volume, issue and pages [CONFIRM: not yet assigned in Crossref].",
     "https://doi.org/10.1016/j.prro.2026.09.009", "Companion to A5. Absent from ORCID."],
    ["C8", "Anderson BM, Brock KK. Artificial intelligence-based image registration and segmentation. In: Wang Y, Qi XS (eds). Artificial Intelligence in Adaptive Radiation Therapy. IOP Publishing, 11/2025, chapter 9, pp. 9-1 to 9-31. ISBN 9780750361194.",
     "https://doi.org/10.1088/978-0-7503-6119-4ch9",
     "Book chapter, published (Crossref record checked 27 Oct 2026). The CV lists it as “Chapter 9: Imaging Registration and Auto-Segmentation” with no publisher or date; the published chapter title is shown here."],
    ["—", "Not eligible: the medRxiv preprint (Castelo et al., 02/2026) until accepted; conference abstracts and presentations; invited talks.", "—",
     "NSF: “unpublished documents not yet submitted for publication, invited lectures, and additional lists of products” are unacceptable."],
]

OPEN_ITEMS = [
    "Position title wording (Clinical Assistant Professor of Medical Physics vs. Assistant Professor, Medical Physics). Use the same wording in all four NSF documents.",
    "Completion dates: Ph.D. 05/2021 (NIH biosketch) or 06/2021 (SciENcv 02/2026); Residency 07/2023 (NIH) or 06/2023 (SciENcv). Check the diploma / residency certificate.",
    "Field of study for the S.M.S. and Ph.D.: plain “Medical Physics” (SciENcv 02/2026) or the NIH biosketch’s “Therapeutic” / “Diagnostic Imaging Medical Physics”.",
    "Residency entry type in the SciENcv Professional Preparation drop-down (Residency / Other / typed “Resident” as in 02/2026).",
    "Department names flagged in section 3: UNC adjunct appointment, MD Anderson graduate research, St. Joseph Hospital fellowship, Georgia Tech research assistantship; and a city/state for the Medical Physics Associate Editor entry.",
    "Whether to list Clinical Resident, Chief Resident, Ad hoc Associate Editor, and the AAPM committee memberships as appointments (section 3b). Recommendation: keep the titled roles, leave committee memberships for Synergistic Activities.",
    "TG-263 membership end date: CV Professional Service says 2021 – Present; CV Major Contributions says 2021–2025.",
    "Products: confirm the proposal topic and the A/B split (section 4). Swap in bench items (4c) as needed; at most five per list.",
    "Volume/issue/pages for A5 (Advances in Radiation Oncology 2026) and C7 (Practical Radiation Oncology 2026); first Keras release containing B5.",
    "ORCID gap: a previous session found Brian’s ORCID record missing works from 2022 onward. A2, A5, B3, B5 (and C6, C7, C8 if used) will not import; enter them manually in SciENcv or add them to ORCID first and re-sync.",
    "Certify in SciENcv from Brian’s own account; the certification date must fall within 12 months of proposal submission. The 02/2026 certified PDF in this repo is the NIH common form, not the NSF document, and cannot be submitted to NSF.",
    "CV discrepancies noticed while verifying citations (not changed here; separate /update-cv task): Open RT Structures date and author list; DicomRTTool paper title/date; DICOM attribute tool title (SOP vs Study Instance UID); COVER-ALL title and “Shan”/“Shah”; Frontiers journal name; elective-nodes paper author list (Dooley J); book chapter title, publisher (IOP) and date (11/2025).",
]

# ----------------------------------------------------------------------------- build

def build():
    doc = Document()
    style_fonts(doc)
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.6))
    footer(doc)

    para(doc, "NSF Biographical Sketch — SciENcv Entry Worksheet", bold=True, size=18, after=2)
    para(doc, f"Brian Mark Anderson, PhD, DABR · University of California, San Diego · prepared {PREPARED}", italic=True, after=10)

    rich(doc, [("Purpose. ", True, False),
               ("NSF requires the Biographical Sketch to be generated in SciENcv (PAPPG NSF 24-1, Chapter II.D.2.h(i); NSPM-33 common form, OMB 3145-0279). "
                "This worksheet lists every value to enter, in the order SciENcv asks for it, with the source document for each value. "
                "It is a proposal for Brian to review: nothing has been entered into SciENcv, and no existing document was modified.", False, False)])
    rich(doc, [("Sources. ", True, False),
               ("BMAnderson CV.docx (saved 30 Sep 2026); Biosketch_BMA.docx (NIH format, 12 Aug 2026); Anderson_CommonFormBiosketch.pdf "
                "(NIH common form generated and certified in SciENcv on 26 Feb 2026; the only document that records start months); "
                "Crossref, PubMed Central, PyPI and GitHub records checked on 27 Oct 2026.", False, False)])
    rich(doc, [("Legend. ", True, False),
               ("[CONFIRM] = the documents disagree or are silent; decide before entering. [DOI NEEDED] = no persistent identifier could be verified; never guess a DOI. "
                "Unflagged values trace directly to a source document.", False, False)])
    rich(doc, [("Form shape. ", True, False),
               ("The NSF common form has Identifying Information, Organization and Location, Professional Preparation, Appointments and Positions, Products, and Certification. "
                "There is no Synergistic Activities section (that is a separate one-page upload) and no page limit, but content beyond what the form asks for can cause the proposal to be returned without review. "
                "Do not enter personal information (home address, phone, etc.).", False, False)], after=4)

    heading(doc, "1. Identifying Information, Organization and Location")
    table(doc, ["Field", "Value to enter", "Source / notes"], IDENTIFYING, [1.8, 4.4, 3.6])

    heading(doc, "2. Professional Preparation (reverse chronological by start date)")
    para(doc, "SciENcv asks, per entry, for organization, location, degree, start date, month/year of receipt, and field of study. "
              "Start months are not in the CV (years only) or the NIH biosketch (completion months only); they come from the SciENcv biosketch certified on 26 Feb 2026 "
              "and should already be stored in Brian’s SciENcv profile. If the new document is created from that existing document, check that they carry over.", after=4)
    table(doc, PREP_HEAD, PREP, [2.2, 1.4, 1.3, 0.75, 0.95, 1.5, 1.7])

    heading(doc, "3. Appointments and Positions (reverse chronological by start date)")
    rich(doc, [("NSF rule (common form text). ", True, False),
               ("“Appointments and positions include any titled academic, professional, or institutional position whether or not remuneration is received, and whether full-time, part-time, or voluntary (including adjunct, visiting, or honorary).” "
                "Domestic and foreign professional appointments outside of the primary organization must be listed for a period up to three years before the proposal submission date. "
                "Both UNC appointments (Clinical Assistant Professor 2023–2025 and Adjunct Assistant Professor 2025–present) fall inside that window: ", False, False),
               ("do not omit either one.", True, False)], after=4)
    heading(doc, "3a. Entries to list", 2)
    table(doc, APPT_HEAD, APPT, [0.65, 0.7, 1.9, 1.8, 1.6, 1.35, 1.8])
    heading(doc, "3b. Other titled roles and memberships (Brian’s call)", 2)
    para(doc, "NSF counts titled positions; plain committee memberships and society memberships are normally left out of this section and described in Synergistic Activities instead.", after=4)
    table(doc, OPT_HEAD, OPT, [0.75, 0.7, 2.6, 2.3, 3.45])

    heading(doc, "4. Products")
    rich(doc, [("Assumption. ", True, False),
               ("The proposal topic was not known when this worksheet was prepared. The split below assumes a proposal on AI / deep learning for radiation oncology "
                "(auto-segmentation, deformable registration, clinical software and data infrastructure). If the proposal is about something else, swap items between 4a, 4b and the bench list in 4c. "
                "NSF allows up to five products in each list; only these ten are used in review.", False, False)])
    rich(doc, [("Entry rules. ", True, False),
               ("Products must be citable and accessible: publications, software, datasets, patents and the like; unpublished documents, invited lectures and “additional lists” are not allowed. "
                "Each entry needs authors (“et al.” allowed after some authors), title, venue with volume/issue/pages where applicable, release date, and a DOI or URL. "
                "In SciENcv, add products from My Bibliography / ORCID where they exist; everything else is typed in manually. Citations can be reordered by dragging.", False, False)], after=4)
    heading(doc, "4a. Products most closely related to the proposed project (max 5)", 2)
    table(doc, PROD_HEAD, RELATED, [0.35, 1.9, 2.3, 1.75, 0.75, 1.5, 1.25])
    heading(doc, "4b. Other significant products (max 5)", 2)
    table(doc, PROD_HEAD, OTHER, [0.35, 1.9, 2.3, 1.75, 0.75, 1.5, 1.25])
    heading(doc, "4c. Bench: verified alternates to swap in", 2)
    table(doc, BENCH_HEAD, BENCH, [0.35, 5.5, 1.9, 2.05])

    heading(doc, "5. Certification")
    para(doc, "The certification text lives inside SciENcv and is not typed in. After the sections above are complete, click Download PDF; SciENcv shows the certification "
              "(“I certify that the information provided is current, accurate, and complete … I am not a party to a malign foreign talent recruitment program …”) and Brian must click CERTIFY "
              "from his own NCBI account. A delegate can draft the document but cannot certify it. The certification date must be within 12 months of the proposal submission date; "
              "the PDF is then uploaded to the proposal in Research.gov.", after=4)

    heading(doc, "6. Open items for Brian")
    for item in OPEN_ITEMS:
        bullet(doc, item)

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
