import sys, json
from openpyxl import load_workbook
orig = load_workbook(sys.argv[1]); o = orig.active
out = load_workbook(sys.argv[2]); w = out.active
data = json.load(open(sys.argv[3], encoding="utf-8"))
N = 171  # rows inserted (detected from table ref growth)
def sty(c):
    return (c.font.name, c.font.sz, c.font.b, c.font.i, c.font.color.rgb if c.font.color else None,
            c.fill.fill_type, c.fill.fgColor.rgb if c.fill.fill_type else None,
            c.border.left.style, c.border.right.style, c.border.top.style, c.border.bottom.style,
            c.alignment.wrap_text, c.alignment.horizontal, c.alignment.vertical, c.number_format)
problems = []
# 1. everything above Table 4's data area (rows 1-51) must be identical except cleared sample values in tables 1-3
sample_rows = set(range(17,22)) | set(range(28,33)) | set(range(38,45))
for r in range(1, 52):
    for c in range(1, 8):
        a, b = o.cell(r, c), w.cell(r, c)
        if sty(a) != sty(b): problems.append(f"style diff {a.coordinate}: {sty(a)} vs {sty(b)}")
        if r not in sample_rows and a.value != b.value: problems.append(f"value diff {a.coordinate}: {a.value!r} vs {b.value!r}")
# 2. rows 57..68 of the original must equal rows 57+N.. of the output (styles; values except table-5 samples)
for r in range(57, o.max_row + 1):
    for c in range(1, 8):
        a, b = o.cell(r, c), w.cell(r + N, c)
        if sty(a) != sty(b): problems.append(f"style diff shifted {a.coordinate}->{b.coordinate}")
        if r not in range(64, 69) and a.value != b.value: problems.append(f"value diff shifted {a.coordinate}->{b.coordinate}: {a.value!r} vs {b.value!r}")
# 3. Table 4 rows: style must match the original first data row (52), column by column
for r in range(52, 52 + 175):
    for c in range(1, 6):
        if sty(w.cell(r, c)) != sty(o.cell(52, c)): problems.append(f"table4 style diff {w.cell(r,c).coordinate}")
# 4. row heights shifted
for idx, rd in o.row_dimensions.items():
    if rd.height and idx >= 57:
        if w.row_dimensions[idx + N].height != rd.height: problems.append(f"row height {idx}->{idx+N}: {rd.height} vs {w.row_dimensions[idx+N].height}")
    elif rd.height and idx < 52 and w.row_dimensions[idx].height != rd.height: problems.append(f"row height {idx}")
# 5. column widths, sheet name, page setup
if o.title != w.title: problems.append("sheet name")
for k, cd in o.column_dimensions.items():
    if w.column_dimensions[k].width != cd.width: problems.append(f"col width {k}")
if (o.page_setup.scale, o.page_setup.orientation, o.page_setup.fitToHeight) != (w.page_setup.scale, w.page_setup.orientation, w.page_setup.fitToHeight): problems.append("page setup")
# 6. content
t4 = [(w.cell(r,1).value, w.cell(r,2).value, w.cell(r,3).value, w.cell(r,4).value, w.cell(r,5).value) for r in range(52, 52+176)]
filled = [t for t in t4 if any(v not in (None, "") for v in t)]
print("Table 4 filled rows:", len(filled), "expected", len(data["table4"]))
print("types:", set(t[0] for t in filled))
print("names all 'Surname, Given' with comma:", all("," in (t[1] or "") for t in filled))
exp = [(r["name"], r["org"]) for r in data["table4"]]
got = [(t[1], t[2]) for t in filled]
print("names/orgs match table_data.json:", exp == got)
print("[CONFIRM] rows in sheet:", sum(1 for t in filled if t[3] and "[CONFIRM]" in t[3]), "expected", sum(1 for r in data["table4"] if r["flag"]))
alph = [c.coordinate for row in w.iter_rows() for c in row if c.value and "Alphaman" in str(c.value) or (c.value and "Samplename" in str(c.value))]
print("leftover sample text:", alph)
print("Table 1:", [(w.cell(r,2).value, w.cell(r,3).value, w.cell(r,4).value) for r in range(17,22) if w.cell(r,3).value])
print("Table 2 non-empty:", [r for r in range(28,33) if any(w.cell(r,c).value for c in range(1,6))])
print("Table 3:", [(w.cell(r,1).value, w.cell(r,2).value, w.cell(r,3).value, w.cell(r,4).value) for r in range(38,45) if w.cell(r,2).value])
print("Table 5:", [(w.cell(r,1).value, w.cell(r,2).value, w.cell(r,3).value, w.cell(r,4).value, w.cell(r,5).value) for r in range(235,240) if w.cell(r,2).value])
print("longest optional:", max((len(t[3] or ""), t[3]) for t in filled))
print("max_row:", w.max_row)
print("PROBLEMS:", len(problems)); print("\n".join(problems[:40]))
