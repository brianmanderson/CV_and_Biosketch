# -*- coding: utf-8 -*-
"""Print the structure of the NSF COA template: sheets, dimensions, merged ranges, every non-empty cell,
data validations, and row heights, so the fill script can be written against the real layout."""
import sys
from openpyxl import load_workbook

wb = load_workbook(sys.argv[1])
print("sheets:", wb.sheetnames)
for ws in wb.worksheets:
    print(f"\n=== {ws.title!r} dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
    print("merged:", [str(r) for r in ws.merged_cells.ranges])
    print("col widths:", {k: v.width for k, v in ws.column_dimensions.items()})
    print("data validations:", [(dv.sqref, dv.type, dv.formula1) for dv in ws.data_validations.dataValidation])
    print("cond formats:", [str(r) for r in ws.conditional_formatting])
    print("freeze:", ws.freeze_panes, " print area:", ws.print_area)
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and str(c.value).strip():
                v = str(c.value).replace("\n", "\\n")
                print(f"{c.coordinate:6s} bold={c.font.b!s:5s} fill={c.fill.fgColor.rgb if c.fill and c.fill.fill_type else None} "
                      f"border_b={c.border.bottom.style} wrap={c.alignment.wrap_text} | {v[:110]}")
    heights = {r: ws.row_dimensions[r].height for r in range(1, ws.max_row + 1) if ws.row_dimensions[r].height}
    print("row heights:", heights)
