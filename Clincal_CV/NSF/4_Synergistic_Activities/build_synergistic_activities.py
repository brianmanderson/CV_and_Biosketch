# -*- coding: utf-8 -*-
"""Build Anderson_BM_NSF_Synergistic_Activities.docx (NSF PAPPG 24-1 II.D.2.h(iv)).

Usage:  python build_synergistic_activities.py <output_dir>

Formatting: US Letter, 1-inch margins, Times New Roman 11 pt, single spacing.
Every number and date below is taken from Clincal_CV/BMAnderson CV.docx
(see README.md in this folder for the line-by-line trace).
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT = "Times New Roman"
SIZE = Pt(11)

HEADER = ("Synergistic Activities – Brian Mark Anderson, PhD, DABR, "
          "Assistant Professor of Medical Physics, University of California, San Diego")

# Each item: (bold lead-in, body text).
ITEMS = [
    ("Open-source software for deep learning in radiation oncology.",
     " Author and maintainer of DicomRTTool, a Python library that converts DICOM "
     "images, RT structures, and dose grids into NumPy/SimpleITK arrays and converts "
     "model predictions back into clinically importable RT structures (Practical "
     "Radiation Oncology, 2021); ~1,700 downloads per month, 100+ GitHub stars, 30+ "
     "forks, and in use across academic and industry groups internationally. Extended "
     "this work into the Keras core library by authoring the reconstruct_patches "
     "operation and ReconstructPatches2D/3D layers (merged 2026), which, because "
     "Keras 3 is backend-agnostic, reach TensorFlow, PyTorch, and JAX users through a "
     "library with ~18 million monthly installs. Also released the DICOM Attribute "
     "Manipulation Tool for frame-of-reference and UID editing (Journal of Applied "
     "Clinical Medical Physics, 2025)."),

    ("Nomenclature standardization: AAPM TG-263 and Open RT Structures.",
     " Member of AAPM Task Group 263 (Standardizing Nomenclatures in Radiation "
     "Oncology) since 2021 and lead author of Open RT Structures, a freely available "
     "implementation of TG-263 that closed the gap between a published nomenclature and "
     "one clinics could actually adopt (International Journal of Radiation Oncology, "
     "Biology, Physics, 2024). Co-authored the TG-263U1 survey of roadblocks to "
     "standardized nomenclature (Journal of Applied Clinical Medical Physics, 2024) and "
     "presented the adoption tool at the AAMD Annual Conference (2023). As Technical "
     "Service Chief for the central nervous system site at UC San Diego "
     "(2025–present), authored the departmental CNS Standards and Guidelines "
     "defining site-specific treatment protocols, critical-structure nomenclature, and "
     "dose constraints."),

    ("Translation of dissertation research into a randomized clinical trial.",
     " Doctoral work at MD Anderson on biomechanical-model deformable image "
     "registration and deep-learning segmentation for liver tumor ablation (Society of "
     "Interventional Radiology Allied Scientist Grant, $40,000, 2019; Medical Physics, "
     "2021) was translated directly into COVER-ALL (NCT04083378), a Phase 2 randomized "
     "controlled trial of software-based versus visual assessment of the minimal "
     "ablative margin at the time of the procedure. The trial established "
     "software-based assessment as superior to visual assessment (The Lancet "
     "Gastroenterology & Hepatology, 2025). The dissertation received MD "
     "Anderson’s Alfred G. Knudson Jr. Outstanding Dissertation Award (2022)."),

    ("Hands-on deep-learning training for medical physicists.",
     " Created and taught “getting started with deep learning” workshops "
     "that take clinical physicists from DICOM to model predictions, delivered at the "
     "Winter Institute of Medical Physics (2020, 2025) and the MD Anderson Image Guided "
     "Cancer Therapy Research Program (2020, 2021). Authored open teaching materials "
     "for the AAPM Practical Big Data Workshop (GitHub/Jupyter setup and TensorFlow "
     "data-pipeline guides, 2021; Early Career Investigator Impact Award) and, as "
     "Teaching Assistant for the 2019 AAPM Summer School, helped create and distribute "
     "its workbooks. Guest lectured on deep learning at Rice University (ELEC/COMP 576, "
     "2019) and the Georgia Institute of Technology (2023), and as Chief Resident at "
     "UC San Diego (2021–2023) coordinated resident education with faculty."),

    ("Service to the scientific community and broadening participation.",
     " Associate Editor of Medical Physics and member of its Board of Associate "
     "Editors (2023–present; ad hoc Associate Editor 2020–2021). Member of the "
     "AAPM Machine Intelligence Subcommittee "
     "(2024–present), the AAPM QUANTEC 2 Working Group (2026–present), and "
     "Medical Physicists for World Benefit (2017–present). Founding member and "
     "domestic liaison of MD Anderson’s International Students Association "
     "(2019–2021), which helps international graduate students with the transition "
     "to the United States; STEM outreach to local students through AAPM WizKids "
     "(2018–2019) and UTHealth Science Night (2017–2019); and Co-Investigator "
     "on AHRQ R18 HS029474 (2024–2025), developing AI-enhanced pre-treatment peer "
     "review to improve patient safety in radiation oncology."),
]


def _force_rfonts(rpr):
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


def set_run_font(run, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = SIZE
    run.font.bold = bold
    run.font.italic = italic
    _force_rfonts(run._element.get_or_add_rPr())


def single_space(pf, after_pt=0):
    pf.space_before = Pt(0)
    pf.space_after = Pt(after_pt)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.widow_control = True


def main(out_dir: Path):
    doc = Document()

    # Page setup: US Letter, 1-inch margins on all sides.
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Inches(1))
    sec.header_distance = Inches(0.5)
    sec.footer_distance = Inches(0.5)

    # Normal style: Times New Roman 11 pt, single spaced, no extra spacing.
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = SIZE
    _force_rfonts(normal.element.get_or_add_rPr())
    single_space(normal.paragraph_format)

    # Header line.
    p = doc.add_paragraph()
    single_space(p.paragraph_format, after_pt=6)
    set_run_font(p.add_run(HEADER), bold=True)

    # Numbered items with a hanging indent. Numbers are typed, not a Word list,
    # so nothing depends on numbering definitions.
    hang = Inches(0.25)
    for n, (lead, body) in enumerate(ITEMS, start=1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        single_space(pf, after_pt=6 if n < len(ITEMS) else 0)
        pf.left_indent = hang
        pf.first_line_indent = -hang
        pf.tab_stops.add_tab_stop(hang, WD_TAB_ALIGNMENT.LEFT)
        set_run_font(p.add_run(f"{n}.\t"))
        set_run_font(p.add_run(lead), bold=True)
        set_run_font(p.add_run(body))

    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / "Anderson_BM_NSF_Synergistic_Activities.docx"
    doc.save(out)
    per_item = [len(l) + len(b) for l, b in ITEMS]
    print(f"wrote {out}")
    print(f"header chars: {len(HEADER)}; body chars: {sum(per_item)}; per item: {per_item}")


if __name__ == "__main__":
    main(Path(sys.argv[1]))
